from flask import Flask, render_template, request, send_file, redirect, url_for
import os
import difflib
from werkzeug.utils import secure_filename
import io
import pandas as pd
from fpdf import FPDF
import nltk
from nltk.corpus import wordnet
import pytesseract
from PIL import Image
import pdfplumber
import docx
import hashlib
from collections import defaultdict

# NLTK setup
try:
    wordnet.synsets('test')
except LookupError:
    nltk.download('wordnet')
    nltk.download('stopwords')
    nltk.download('omw-1.4')

app = Flask(__name__)
UPLOAD_FOLDER = 'uploads'
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER
os.makedirs(UPLOAD_FOLDER, exist_ok=True)

latest_results = []
detailed_matches = {}

@app.route('/')
def index():
    return render_template('index.html', results=latest_results)

@app.route('/compare', methods=['POST'])
def compare():
    global latest_results, detailed_matches
    latest_results = []
    detailed_matches = {}

    for f in os.listdir(app.config['UPLOAD_FOLDER']):
        os.remove(os.path.join(app.config['UPLOAD_FOLDER'], f))

    original_file = request.files['original']
    original_filename = secure_filename ( original_file.filename )
    original_path = os.path.join ( app.config['UPLOAD_FOLDER'], "original_" + original_filename )
    original_file.save(original_path)
    original_text = read_file(original_path)

    student_files = request.files.getlist('students')
    student_text_dict = {}

    for student_file in student_files:
        filename = secure_filename(student_file.filename)
        if not filename.lower().endswith(('.txt', '.docx', '.pdf', '.png', '.jpg', '.jpeg')):
            continue
        filepath = os.path.join(app.config['UPLOAD_FOLDER'], filename)
        student_file.save(filepath)
        student_text_dict[filename] = read_file(filepath)

    duplicates = detect_duplicates(student_text_dict)

    for filename, student_text in student_text_dict.items():
        similarity = calculate_similarity(original_text, student_text)
        verdict = interpret_similarity(similarity)
        copied_from = duplicates[filename]
        latest_results.append((filename, similarity, verdict, copied_from))

        # Store detailed diff for on-demand viewing
        diff_html = generate_diff(original_text, student_text)
        detailed_matches[filename] = diff_html

    return redirect(url_for('index'))

@app.route('/details/<filename>')
def details(filename):
    student_path = os.path.join(app.config['UPLOAD_FOLDER'], filename)

    # Find the original file
    original_file = None
    for f in os.listdir(app.config['UPLOAD_FOLDER']):
        if f.startswith("original_"):
            original_file = os.path.join(app.config['UPLOAD_FOLDER'], f)
            break

    if not original_file:
        return "Original file not found.", 404

    with open(original_file, 'r', encoding='utf-8') as f:
        original_text = f.read()

    with open(student_path, 'r', encoding='utf-8') as f:
        student_text = f.read()

    # Use difflib to highlight matching sections
    matcher = difflib.SequenceMatcher(None, original_text, student_text)

    highlighted_student = ""
    for opcode, i1, i2, j1, j2 in matcher.get_opcodes():
        student_segment = student_text[j1:j2]
        if opcode == 'equal' and (i2 - i1) > 10:  # Only highlight significant matches
            highlighted_student += f"<mark style='background-color:#ffb3b3'>{student_segment}</mark>"
        else:
            highlighted_student += student_segment

    return render_template('details.html',
                           filename=filename,
                           original_text=original_text,
                           highlighted_student=highlighted_student)

@app.route('/download-report')
def download_report():
    format = request.args.get('format', 'excel')
    if not latest_results:
        return "No report available.", 400

    if format == 'excel':
        df = pd.DataFrame(latest_results, columns=["Student File", "Similarity (%)", "Verdict", "Copied From"])
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, index=False)
        output.seek(0)
        return send_file(output, as_attachment=True, download_name="similarity_report.xlsx",
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    elif format == 'pdf':
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", size=20)
        pdf.cell(0, 12, txt="SIMILARITY REPORT", ln=True, align='C')
        pdf.ln(8)
        pdf.set_font("Arial", "B", size=10)
        pdf.cell(50, 8, txt="Filename", border=1)
        pdf.cell(30, 8, txt="Similarity", border=1)
        pdf.cell(35, 8, txt="Verdict", border=1)
        pdf.cell(65, 8, txt="Copied From", border=1)
        pdf.ln()

        pdf.set_font("Arial", size=9)
        for filename, similarity, verdict, copied_from in latest_results:
            pdf.cell(50, 8, txt=filename[:25], border=1)
            pdf.cell(30, 8, txt=str(similarity) + "%", border=1)
            pdf.cell(35, 8, txt=verdict, border=1)
            pdf.multi_cell ( 65, 8, txt=copied_from, border=1 )


        pdf_output = io.BytesIO()
        pdf_bytes = pdf.output(dest='S').encode('latin1')
        pdf_output.write(pdf_bytes)
        pdf_output.seek(0)
        return send_file(pdf_output, as_attachment=True, download_name="similarity_report.pdf",
                         mimetype='application/pdf')
    else:
        return "Invalid format requested.", 400

# ========== Helper Functions ==========

def read_file(path):
    try:
        if path.lower().endswith(".txt"):
            with open(path, 'r', encoding="utf-8") as f:
                return f.read()
        elif path.lower().endswith(".docx"):
            doc = docx.Document(path)
            return "\n".join([para.text for para in doc.paragraphs])
        elif path.lower().endswith(".pdf"):
            text = ""
            with pdfplumber.open(path) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
            if not text.strip():
                with pdfplumber.open(path) as pdf:
                    for page in pdf.pages:
                        image = page.to_image(resolution=300).original
                        text += pytesseract.image_to_string(image)
            return text
        elif path.lower().endswith(('.png', '.jpg', '.jpeg')):
            image = Image.open(path)
            return pytesseract.image_to_string(image)
    except Exception as e:
        return ""

def synonym_normalize(text):
    words = text.split()
    normalized = []
    for word in words:
        synsets = wordnet.synsets(word)
        if synsets:
            lemma = synsets[0].lemmas()[0].name().replace('_', ' ')
            normalized.append(lemma.lower())
        else:
            normalized.append(word.lower())
    return ' '.join(normalized)

def calculate_similarity(text1, text2):
    text1 = synonym_normalize(text1)
    text2 = synonym_normalize(text2)
    sequence = difflib.SequenceMatcher(None, text1, text2)
    return round(sequence.ratio() * 100, 2)

def interpret_similarity(similarity):
    if similarity >= 80:
        return "Definitely Copied"
    elif similarity >= 50:
        return "Somewhat Copied"
    elif similarity >= 20:
        return "Slightly Similar"
    else:
        return "Unique Content"

def get_md5(text):
    return hashlib.md5(text.encode('utf-8')).hexdigest()

def detect_duplicates(student_text_dict):
    hash_map = defaultdict(list)
    for filename, text in student_text_dict.items():
        file_hash = get_md5(text)
        hash_map[file_hash].append(filename)
    duplicates = {}
    for files in hash_map.values():
        for file in files:
            others = [f for f in files if f != file]
            duplicates[file] = ', '.join(others) if others else '-'
    return duplicates

def generate_diff(text1, text2):
    differ = difflib.HtmlDiff()
    return differ.make_table(text1.splitlines(), text2.splitlines(), "Original", "Student", context=True)

if __name__ == '__main__':
    app.run(debug=True)
